"""Reproducible, synthetic QA assets. Never writes to the live platform."""
from pathlib import Path
import csv
import hashlib
import json
import math
import struct
import wave
from xml.sax.saxutils import escape

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'tests/fixtures/qa-round2'
OUT.mkdir(parents=True, exist_ok=True)
FONT = ROOT / 'public/fonts/DejaVuSans.ttf'
if not FONT.exists():
    FONT = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
pdfmetrics.registerFont(TTFont('QASans', str(FONT)))
style = ParagraphStyle('qa', fontName='QASans', fontSize=9, leading=13, textColor=colors.HexColor('#123b34'))
head = ParagraphStyle('heading', parent=style, fontSize=15, leading=20, spaceAfter=12)

FORMATS = ['Selección directa', 'Completar', 'Relacionar', 'Ordenar', 'Caso práctico']
COLUMNS = ['pregunta','opcion_a','opcion_b','opcion_c','opcion_d','correcta','explicacion','formato','contexto','fuente','mezclar_alternativas']
# Each correct option begins first; rotate it below so keys A, B, C and D are tested.
BANK = [
('Selección directa','¿Qué acción permite comprobar si la cantidad de agua influye en la germinación?', ['Cambiar solo el agua y mantener las otras condiciones iguales.','Cambiar agua, luz y suelo a la vez.','Concluir antes de observar.','Omitir el registro de las semillas.'], 'Una sola variable modificada permite comparar el efecto del agua.'),
('Selección directa','¿Qué instrumento permite medir el volumen de agua?', ['Una probeta graduada.','Un termómetro.','Un cronómetro.','Una balanza sin recipiente.'], 'La probeta graduada permite leer el volumen del líquido.'),
('Selección directa','¿Qué dato representa una observación medible?', ['La planta creció 3 centímetros.','La planta parece feliz.','La planta debería crecer.','La planta es la mejor.'], 'Los centímetros expresan una medida verificable.'),
('Selección directa','¿Qué evidencia ayuda a comparar dos experimentos?', ['Registros obtenidos con el mismo procedimiento.','Opiniones sin registros.','Solo el nombre de los participantes.','Resultados elegidos para confirmar una idea.'], 'Un procedimiento comparable y registros completos permiten contrastar resultados.'),
('Completar','La cantidad de agua que se cambia entre dos grupos es la variable ______.', ['independiente','dependiente','inexistente','de conclusión'], 'La variable independiente es la que se modifica deliberadamente.'),
('Completar','La longitud de la planta, medida cada día, constituye un ______.', ['dato','supuesto sin comprobar','nombre de grupo','comprobante'], 'La longitud medida es un dato registrado del experimento.'),
('Completar','Para comunicar una medición de longitud se indica el número y la ______.', ['unidad','contraseña','opinión','clave del estudiante'], 'Una medición necesita valor y unidad para interpretarse.'),
('Completar','Una conclusión del experimento debe apoyarse en la ______ registrada.', ['evidencia','preferencia','predicción sin comprobar','lista de asistentes'], 'La conclusión utiliza la evidencia obtenida y registrada.'),
('Relacionar','Relaciona instrumento y magnitud.\nLista 1: 1. Termómetro. 2. Cronómetro.\nLista 2: a. Tiempo. b. Temperatura.', ['1-b; 2-a','1-a; 2-b','1-a; 2-a','1-b; 2-b'], 'El termómetro mide temperatura y el cronómetro mide tiempo.'),
('Relacionar','Relaciona acción y finalidad.\nLista 1: 1. Observar. 2. Registrar.\nLista 2: a. Conservar datos. b. Examinar cambios.', ['1-b; 2-a','1-a; 2-b','1-a; 2-a','1-b; 2-b'], 'Observar examina cambios; registrar conserva los datos.'),
('Relacionar','Relaciona unidad y magnitud.\nLista 1: 1. Centímetro. 2. Segundo.\nLista 2: a. Tiempo. b. Longitud.', ['1-b; 2-a','1-a; 2-b','1-a; 2-a','1-b; 2-b'], 'El centímetro expresa longitud y el segundo expresa tiempo.'),
('Relacionar','Relaciona registro y propósito.\nLista 1: 1. Tabla de datos. 2. Gráfico.\nLista 2: a. Visualizar una comparación. b. Conservar valores organizados.', ['1-b; 2-a','1-a; 2-b','1-a; 2-a','1-b; 2-b'], 'La tabla conserva valores y el gráfico facilita su comparación.'),
('Ordenar','Ordena una investigación breve.\n1. Comunicar resultados.\n2. Observar y registrar.\n3. Preparar las condiciones.', ['3 → 2 → 1','1 → 2 → 3','2 → 1 → 3','3 → 1 → 2'], 'Primero se prepara, después se observa y finalmente se comunica.'),
('Ordenar','Ordena una medición.\n1. Anotar el valor y la unidad.\n2. Leer la escala.\n3. Elegir el instrumento adecuado.', ['3 → 2 → 1','1 → 2 → 3','2 → 1 → 3','3 → 1 → 2'], 'Elegir el instrumento precede a leerlo y registrar el resultado.'),
('Ordenar','Ordena el análisis de una tabla.\n1. Elaborar una conclusión.\n2. Comparar los valores.\n3. Revisar los registros completos.', ['3 → 2 → 1','1 → 2 → 3','2 → 1 → 3','3 → 1 → 2'], 'Revisar los datos permite compararlos antes de concluir.'),
('Ordenar','Ordena la revisión de un experimento.\n1. Repetir con condiciones controladas.\n2. Identificar el cambio inesperado.\n3. Revisar el procedimiento registrado.', ['3 → 2 → 1','1 → 2 → 3','2 → 1 → 3','3 → 1 → 2'], 'Se revisa el registro, se identifica el problema y se repite controlando condiciones.'),
('Caso práctico','¿Qué decisión permite comparar el efecto del agua?', ['Mantener iguales la luz, las semillas y el suelo.','Cambiar también el tipo de semilla.','Omitir el grupo de comparación.','Eliminar las mediciones que no coincidan.'], 'Para comparar el efecto del agua deben controlarse las otras condiciones.'),
('Caso práctico','¿Qué debe hacerse antes de concluir que el tratamiento funciona?', ['Comparar los registros de ambos grupos.','Elegir el resultado más conveniente.','Usar solo el primer día.','Cambiar la pregunta sin registrar el cambio.'], 'Comparar todos los registros permite fundamentar la conclusión.'),
('Caso práctico','¿Cómo debe tratarse una medición inesperada?', ['Revisar el instrumento y conservar el registro de la revisión.','Borrarla sin anotarlo.','Sustituirla por la predicción.','Copiar el valor de otro equipo.'], 'La revisión transparente conserva la evidencia y permite explicar errores.'),
('Caso práctico','¿Qué registro mejora la posibilidad de repetir la experiencia?', ['Materiales, procedimiento, condiciones y datos obtenidos.','Solo la conclusión final.','Solo los nombres del equipo.','Solo la respuesta preferida.'], 'Un registro del procedimiento y las condiciones permite repetir la experiencia.'),
]

manifest = {'label':'TUTOSEBAS · segunda ronda de QA', 'synthetic':True, 'period':'2026-2027', 'blocks':[], 'files':[]}
for number, (subject, topic) in enumerate([
    ('Didáctica de Ciencias Naturales QA', 'Tema 1 · Observación'),
    ('Didáctica de Ciencias Naturales QA', 'Tema 2 · Experimentos'),
    ('Investigación Educativa QA', 'Tema 1 · Registro'),
    ('Investigación Educativa QA', 'Tema 2 · Análisis'),
], 1):
    slug = f'bloque-{number:02d}'
    questions = []
    for index, (fmt, prompt, options, explanation) in enumerate(BANK):
        correct = index % 4
        rotated = options[-correct:] + options[:-correct] if correct else list(options)
        question = {'prompt':f'Experiencia {number}: {prompt}', 'options':rotated, 'correctIndex':correct,
                    'format':fmt, 'caseContext':f'Un equipo del taller {number} compara dos grupos de semillas. Cambia la cantidad de agua y registra diariamente su longitud.\nEl registro conserva las condiciones de luz y suelo.' if fmt == 'Caso práctico' else '',
                    'explanation':explanation, 'source':f'Material de prueba QA · Experiencia {number} · Sección {index+1}', 'shuffleOptions':fmt not in ['Ordenar','Relacionar']}
        questions.append(question)
    lines = []
    for i, q in enumerate(questions, 1):
        lines.extend([f'Pregunta {i}: {q["prompt"] if not q["caseContext"] else ""}', f'Formato: {q["format"]}'])
        if q['caseContext']: lines.extend([f'Contexto: {q["caseContext"]}',f'Enunciado: {q["prompt"]}'])
        lines.extend([f'Alternativa {chr(65+j)}: {v}' for j,v in enumerate(q['options'])])
        lines.extend([f'Correcta: {chr(65+q["correctIndex"])}',f'Explicación: {q["explanation"]}',f'Fuente: {q["source"]}',f'Mezclar alternativas: {"Sí" if q["shuffleOptions"] else "No"}', ''])
    text = '\n'.join(lines)
    (OUT/f'{slug}.txt').write_text(text, encoding='utf-8')
    doc = Document()
    for line in text.split('\n'): doc.add_paragraph(line)
    doc.save(OUT/f'{slug}.docx')
    table = [COLUMNS] + [[q['prompt'],*q['options'],chr(65+q['correctIndex']),q['explanation'],q['format'],q['caseContext'],q['source'],'Sí' if q['shuffleOptions'] else 'No'] for q in questions]
    with (OUT/f'{slug}.csv').open('w',newline='',encoding='utf-8-sig') as file:
        writer=csv.writer(file,delimiter=';',quoting=csv.QUOTE_ALL);writer.writerows(table)
    wb=Workbook();ws=wb.active;ws.title='Preguntas'
    for row in table: ws.append(row)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='064A3D')
    for col in 'ABCDEFGHIJK': ws.column_dimensions[col].width=34
    for row in ws:
        for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
    wb.save(OUT/f'{slug}.xlsx')
    stories=[]
    # No page numbers or repeating labels in the extraction stream.
    for group in text.strip().split('\n\n'):
        stories.append(KeepTogether([Paragraph(escape(line).replace('\n','<br/>'), style) for line in group.split('\n')]))
        stories.append(Spacer(1,12))
    SimpleDocTemplate(str(OUT/f'{slug}.pdf'),pagesize=A4,leftMargin=38,rightMargin=38,topMargin=32,bottomMargin=32).build(stories)
    manifest['blocks'].append({'id':slug,'subject':subject,'topic':topic,'count':20,'formatCounts':{f:4 for f in FORMATS},'questions':questions})

SimpleDocTemplate(str(OUT/'material-lectura.pdf'),pagesize=A4).build([
    Paragraph('TUTOSEBAS · Material de prueba',head),
    Paragraph('Observación, medición y registro de experimentos. Este material es sintético para comprobar archivos protegidos e importación; no es un currículo oficial.',style),Spacer(1,14),
    Paragraph('Una medición incluye valor y unidad. El termómetro mide temperatura, el cronómetro tiempo y la probeta volumen. Una comparación controla las condiciones y modifica una variable.',style),Spacer(1,14),
    Paragraph('Preparar condiciones → observar y registrar → comunicar resultados. Las conclusiones se apoyan en registros completos. Una medición inesperada se revisa y se documenta.',style),
])
im=Image.new('RGB',(1000,460),'#f4faf7');drawing=ImageDraw.Draw(im)
f=ImageFont.truetype(str(FONT),28)
drawing.text((35,25),'QA · Registro experimental (datos ficticios)',fill='#064a3d',font=f)
for index,value in enumerate([3,6,9,12]):
    x=100+index*220;top=380-value*20
    drawing.rectangle((x,top,x+90,380),fill='#147662')
    drawing.text((x,385),f'Día {index+1}',fill='#123b34',font=f)
    drawing.text((x,top-35),f'{value} cm',fill='#123b34',font=f)
im.save(OUT/'infografia.png')
with wave.open(str(OUT/'audio-prueba.wav'),'wb') as audio:
    audio.setnchannels(1);audio.setsampwidth(2);audio.setframerate(16000)
    audio.writeframes(b''.join(struct.pack('<h',round(4000*math.sin(2*math.pi*440*i/16000))) for i in range(16000*3)))
doc=Document();doc.add_heading('QA · Guía de observación',0);doc.add_paragraph('Material sintético. Registrar valor, unidad, fecha y condiciones.');doc.save(OUT/'material-word.docx')
# Real image-only PDF and malformed files for negative cases.
im.convert('RGB').save(OUT/'escaneo-sin-texto.pdf','PDF')
(OUT/'archivo-falso.pdf').write_bytes(b'NO ES UN PDF, prueba de rechazo')
(OUT/'word-danado.docx').write_bytes(b'PK\x03\x04archivo incompleto')
bad=table.copy();bad=[list(row) for row in bad];bad[1][5]='';bad[2][8]='';bad[3][7]='Formato desconocido'
with (OUT/'bloque-errores.csv').open('w',newline='',encoding='utf-8-sig') as file: csv.writer(file,delimiter=';',quoting=csv.QUOTE_ALL).writerows(bad)
for file in sorted(OUT.iterdir()):
    if file.name!='manifest.json': manifest['files'].append({'name':file.name,'size':file.stat().st_size,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()})
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'files':len(manifest['files']),'blocks':4,'questions':80,'path':str(OUT)},ensure_ascii=False))
